import io
import fitz  # PyMuPDF
import docx
from typing import Tuple, List
from app.config import settings
from app.models.schemas import ParsedResume
from app.parser.ocr_detector import detect_scanned_pdf
from app.parser.section_parser import parse_sections, extract_contact_info


class ResumeParsingError(Exception):
    """Base exception for resume parsing errors."""
    pass


class EmptyFileError(ResumeParsingError):
    pass


class OversizedFileError(ResumeParsingError):
    pass


class InvalidFileType(ResumeParsingError):
    pass


class EncryptedPDFError(ResumeParsingError):
    pass


def validate_file_signature(content: bytes, filename: str) -> str:
    """Validates magic bytes against filename extension."""
    lower_name = filename.lower()
    
    if len(content) == 0:
        raise EmptyFileError("Uploaded file is empty (0 bytes).")

    if len(content) > settings.max_upload_size_bytes:
        limit_mb = settings.max_upload_size_bytes / (1024 * 1024)
        raise OversizedFileError(f"File size exceeds the {limit_mb:.0f} MB upload limit.")

    if lower_name.endswith(".pdf"):
        if not content.startswith(b"%PDF-"):
            raise InvalidFileType("File has .pdf extension but does not contain a valid PDF signature.")
        return "pdf"
    
    elif lower_name.endswith(".docx"):
        # DOCX is a ZIP file starting with PK\x03\x04
        if not content.startswith(b"PK\x03\x04"):
            raise InvalidFileType("File has .docx extension but does not contain a valid DOCX signature.")
        return "docx"
    
    else:
        raise InvalidFileType("Unsupported file type. Only PDF and DOCX files are accepted.")


def extract_text_from_pdf(content: bytes, filename: str) -> ParsedResume:
    """Extracts text and structural metadata from a PDF file using PyMuPDF."""
    warnings: List[str] = []
    doc = None
    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception as e:
        raise ResumeParsingError(f"Failed to open PDF document. File may be corrupted: {str(e)}")

    try:
        if doc.is_encrypted:
            raise EncryptedPDFError("The PDF document is password-protected or encrypted. Please provide an unencrypted version.")

        page_count = len(doc)
        if page_count == 0:
            raise EmptyFileError("PDF document has 0 pages.")

        is_scanned, ocr_rec, ocr_msg = detect_scanned_pdf(doc, settings.min_scanned_char_threshold)
        if is_scanned:
            warnings.append(ocr_msg)

        extracted_text_blocks: List[str] = []
        has_tables_or_columns = False

        for page_idx, page in enumerate(doc):
            # Extract blocks: (x0, y0, x1, y1, text, block_no, block_type)
            # block_type 0 is text, 1 is image
            blocks = page.get_text("blocks")
            # Sort vertically by top coordinate (y0) to preserve reading order
            blocks = sorted(blocks, key=lambda b: (b[1], b[0]))
            
            # Simple heuristic for multi-column layout
            # If two text blocks share similar y0 ranges but have significantly separated x coordinates
            for i in range(len(blocks) - 1):
                b1, b2 = blocks[i], blocks[i+1]
                if b1[6] == 0 and b2[6] == 0:  # text blocks
                    if abs(b1[1] - b2[1]) < 10 and abs(b2[0] - b1[2]) > 50:
                        has_tables_or_columns = True

            page_text = "\n".join([b[4].strip() for b in blocks if b[6] == 0 and b[4].strip()])
            if page_text:
                extracted_text_blocks.append(page_text)

        full_raw_text = "\n\n".join(extracted_text_blocks).strip()

        if has_tables_or_columns:
            warnings.append(
                "Possible multi-column or tabular layout detected in PDF. "
                "Some older ATS parsers may read across columns, scrambling content."
            )

        if len(full_raw_text.split()) < 30 and not is_scanned:
            warnings.append("Very low word count detected. Make sure all resume content is visible and unflattened.")

        contact_info = extract_contact_info(full_raw_text)
        sections = parse_sections(full_raw_text)
        word_count = len(full_raw_text.split())

        return ParsedResume(
            filename=filename,
            file_type="pdf",
            page_count=page_count,
            word_count=word_count,
            raw_text=full_raw_text,
            contact_info=contact_info,
            sections=sections,
            is_scanned_pdf=is_scanned,
            ocr_recommended=ocr_rec,
            warnings=warnings
        )
    finally:
        if doc is not None:
            doc.close()


def extract_text_from_docx(content: bytes, filename: str) -> ParsedResume:
    """Extracts text and structure from a DOCX file using python-docx."""
    warnings: List[str] = []
    try:
        doc_stream = io.BytesIO(content)
        doc = docx.Document(doc_stream)
    except Exception as e:
        raise ResumeParsingError(f"Failed to open DOCX document. File may be corrupted or not a valid Word document: {str(e)}")

    paragraphs_text: List[str] = []
    
    # Extract body paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        # If heading style
        if p.style and "Heading" in p.style.name:
            paragraphs_text.append(f"\n{txt}\n")
        elif p.style and "List" in p.style.name:
            paragraphs_text.append(f"• {txt}")
        else:
            paragraphs_text.append(txt)

    # Extract tables if any
    if doc.tables:
        warnings.append(
            "Document contains tables. Some ATS systems struggle to extract text sequentially from table cells. "
            "Consider standard paragraph and bullet formatting."
        )
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                # Deduplicate contiguous identical cells from merged columns
                deduped_cells = []
                for cell_txt in row_cells:
                    if not deduped_cells or cell_txt != deduped_cells[-1]:
                        deduped_cells.append(cell_txt)
                if deduped_cells:
                    paragraphs_text.append(" | ".join(deduped_cells))

    full_raw_text = "\n".join(paragraphs_text).strip()
    
    if len(full_raw_text) == 0:
        raise EmptyFileError("DOCX document contains no readable text.")

    contact_info = extract_contact_info(full_raw_text)
    sections = parse_sections(full_raw_text)
    word_count = len(full_raw_text.split())

    return ParsedResume(
        filename=filename,
        file_type="docx",
        page_count=max(1, word_count // 450),  # Heuristic estimation for docx
        word_count=word_count,
        raw_text=full_raw_text,
        contact_info=contact_info,
        sections=sections,
        is_scanned_pdf=False,
        ocr_recommended=False,
        warnings=warnings
    )


def parse_resume_bytes(content: bytes, filename: str) -> ParsedResume:
    """
    Validates file signature, extracts content, and returns structured ParsedResume.
    Does not persist content to disk (zero retention).
    """
    file_type = validate_file_signature(content, filename)
    
    if file_type == "pdf":
        return extract_text_from_pdf(content, filename)
    elif file_type == "docx":
        return extract_text_from_docx(content, filename)
    else:
        raise InvalidFileType("Unsupported format.")
