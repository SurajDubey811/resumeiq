import io
import re
from typing import List, Dict, Tuple, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from app.parser.section_parser import parse_sections, extract_contact_info, is_header_or_date_line


def apply_accepted_changes(original_text: str, accepted_changes: List[Dict[str, str]]) -> str:
    """Substitutes accepted suggestions into the original resume text."""
    updated = original_text
    for change in accepted_changes:
        orig = change.get("original_text", "").strip()
        sugg = change.get("suggested_text", "").strip()
        if orig and sugg:
            if orig in updated:
                updated = updated.replace(orig, sugg)
            else:
                # Fallback fuzzy match (normalized whitespace)
                orig_norm = " ".join(orig.split())
                lines = updated.splitlines()
                replaced_lines = []
                for line in lines:
                    line_norm = " ".join(line.split())
                    if orig_norm in line_norm:
                        replaced_lines.append(line_norm.replace(orig_norm, sugg))
                    else:
                        replaced_lines.append(line)
                updated = "\n".join(replaced_lines)
    return updated


def generate_updated_docx(updated_text: str, base_name: str) -> bytes:
    """Generates a cleanly formatted ATS-friendly DOCX document."""
    doc = docx.Document()

    # Set 0.75-inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    sections_list = parse_sections(updated_text)
    contact_info = extract_contact_info(updated_text)

    # First line is typically the candidate's name
    lines = [l.strip() for l in updated_text.splitlines() if l.strip()]
    candidate_name = lines[0] if lines else "Candidate Name"
    # If the first line is very long, shorten to reasonable name
    if len(candidate_name) > 40:
        candidate_name = candidate_name.split("|")[0].strip()

    # Title: Name
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(2)
    title_p.paragraph_format.space_before = Pt(0)
    run_name = title_p.add_run(candidate_name)
    run_name.font.size = Pt(18)
    run_name.font.bold = True
    run_name.font.color.rgb = RGBColor(15, 23, 42)  # Slate-900

    # Contact line
    contact_parts = []
    if contact_info.email:
        contact_parts.append(contact_info.email)
    if contact_info.phone:
        contact_parts.append(contact_info.phone)
    if contact_info.location:
        contact_parts.append(contact_info.location)
    if contact_info.linkedin:
        contact_parts.append(contact_info.linkedin)
    if contact_info.github:
        contact_parts.append(contact_info.github)

    if contact_parts:
        contact_p = doc.add_paragraph(" | ".join(contact_parts))
        contact_p.paragraph_format.space_after = Pt(12)
        contact_p.runs[0].font.size = Pt(9.5)
        contact_p.runs[0].font.color.rgb = RGBColor(100, 116, 139)  # Slate-500

    # Sections
    for sec in sections_list:
        if sec.name == "contact_header":
            continue

        # Section Heading
        heading = doc.add_paragraph()
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(4)
        run_h = heading.add_run(sec.raw_title.upper())
        run_h.font.size = Pt(11)
        run_h.font.bold = True
        run_h.font.color.rgb = RGBColor(30, 41, 59)

        # Bullets or Body Paragraphs
        if sec.bullets:
            for b in sec.bullets:
                bp = doc.add_paragraph(style="List Bullet")
                bp.paragraph_format.space_after = Pt(2)
                bp.paragraph_format.space_before = Pt(1)
                r = bp.add_run(b)
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(51, 65, 85)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(sec.content)
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(51, 65, 85)

    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def generate_updated_pdf(updated_text: str, base_name: str) -> bytes:
    """Generates an ATS-compliant PDF document using ReportLab."""
    stream = io.BytesIO()
    doc = SimpleDocTemplate(
        stream,
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()

    # Custom styles
    name_style = ParagraphStyle(
        'ResumeName',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=3,
    )
    contact_style = ParagraphStyle(
        'ResumeContact',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#64748b'),
        spaceAfter=12,
    )
    heading_style = ParagraphStyle(
        'ResumeHeading',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#1e293b'),
        spaceBefore=10,
        spaceAfter=3,
    )
    bullet_style = ParagraphStyle(
        'ResumeBullet',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor('#334155'),
        leftIndent=14,
        firstLineIndent=-10,
        spaceAfter=3,
    )
    body_style = ParagraphStyle(
        'ResumeBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor('#334155'),
        spaceAfter=4,
    )

    story = []

    lines = [l.strip() for l in updated_text.splitlines() if l.strip()]
    candidate_name = lines[0] if lines else "Candidate Name"
    if len(candidate_name) > 40:
        candidate_name = candidate_name.split("|")[0].strip()

    story.append(Paragraph(candidate_name, name_style))

    contact_info = extract_contact_info(updated_text)
    contact_parts = []
    if contact_info.email:
        contact_parts.append(contact_info.email)
    if contact_info.phone:
        contact_parts.append(contact_info.phone)
    if contact_info.location:
        contact_parts.append(contact_info.location)
    if contact_info.linkedin:
        contact_parts.append(contact_info.linkedin)
    if contact_info.github:
        contact_parts.append(contact_info.github)

    if contact_parts:
        story.append(Paragraph(" &bull; ".join(contact_parts), contact_style))

    sections_list = parse_sections(updated_text)

    for sec in sections_list:
        if sec.name == "contact_header":
            continue

        story.append(Paragraph(sec.raw_title.upper(), heading_style))
        story.append(HRFlowable(width="100%", thickness=0.75, color=colors.HexColor('#cbd5e1'), spaceAfter=5))

        if sec.bullets:
            for b in sec.bullets:
                clean_bullet = b.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(f"&bull; {clean_bullet}", bullet_style))
        else:
            clean_content = sec.content.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(clean_content, body_style))

    doc.build(story)
    return stream.getvalue()


def export_resume(
    original_text: str,
    file_type: str,
    accepted_changes: List[Dict[str, str]],
    filename: Optional[str] = "resume",
    target_format: Optional[str] = None
) -> Tuple[bytes, str, str]:
    """
    Exports updated resume in the user's desired format (PDF or DOCX).
    Defaults to matching the uploaded file's format.
    """
    updated_text = apply_accepted_changes(original_text, accepted_changes)
    
    # Clean base name
    clean_base = re.sub(r'\.(pdf|docx)$', '', filename or "resume", flags=re.IGNORECASE)
    
    # Determine format: requested format or original file_type
    fmt = (target_format or file_type or "pdf").lower().strip('.')
    if fmt not in ("pdf", "docx"):
        fmt = "docx" if "docx" in (filename or "").lower() else "pdf"

    if fmt == "docx":
        data = generate_updated_docx(updated_text, clean_base)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        download_name = f"{clean_base}_ATS_Optimized.docx"
    else:
        data = generate_updated_pdf(updated_text, clean_base)
        media_type = "application/pdf"
        download_name = f"{clean_base}_ATS_Optimized.pdf"

    return data, media_type, download_name
