import io
import fitz
import docx
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.parser.extractor import (
    validate_file_signature,
    parse_resume_bytes,
    EmptyFileError,
    OversizedFileError,
    InvalidFileType,
    EncryptedPDFError,
    ResumeParsingError,
)
from app.parser.section_parser import parse_sections, extract_contact_info

client = TestClient(app)


def create_sample_pdf(text: str = None, is_scanned: bool = False) -> bytes:
    """Helper to generate a real, in-memory PDF for testing."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    
    if is_scanned:
        # Create an image-like pixmap and insert it without text
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 100, 100), 0)
        pix.clear_with(200)
        img_bytes = pix.tobytes("png")
        page.insert_image(fitz.Rect(50, 50, 400, 600), stream=img_bytes)
    else:
        content = text or (
            "Jane Doe\n"
            "jane.doe@example.com | (555) 123-4567 | linkedin.com/in/janedoe | github.com/janedoe\n"
            "San Francisco, CA\n\n"
            "Professional Summary\n"
            "Senior Software Engineer with 7+ years of experience designing cloud platforms.\n\n"
            "Technical Skills\n"
            "Python, FastAPI, Docker, Kubernetes, AWS, PostgreSQL, React, TypeScript\n\n"
            "Work Experience\n"
            "Senior Backend Engineer - TechCorp (2020 - Present)\n"
            "• Spearheaded migration of legacy services to FastAPI, boosting throughput by 42%.\n"
            "• Designed resilient distributed event pipeline processing 10M daily events.\n"
            "• Mentored 5 junior engineers and established automated CI/CD testing pipelines.\n\n"
            "Education\n"
            "Bachelor of Science in Computer Science - University of California, Berkeley (2016)\n"
        )
        page.insert_text((50, 72), content, fontsize=11)
        
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx(include_table: bool = False) -> bytes:
    """Helper to generate an in-memory DOCX for testing."""
    doc = docx.Document()
    doc.add_heading("Jane Doe", level=0)
    doc.add_paragraph("jane.doe@example.com | (555) 123-4567 | linkedin.com/in/janedoe")
    
    doc.add_heading("Summary", level=1)
    doc.add_paragraph("Full stack engineer with a passion for high performance systems.")
    
    doc.add_heading("Work Experience", level=1)
    doc.add_paragraph("Lead Developer at Acme Inc. (2021 - Present)")
    p1 = doc.add_paragraph("Built scalable REST APIs reducing latency by 35%.", style="List Bullet")
    p2 = doc.add_paragraph("Automated testing suites achieving 94% test coverage.", style="List Bullet")
    
    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.S. in Computer Science - MIT (2018)")
    
    if include_table:
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Skill"
        table.cell(0, 1).text = "Proficiency"
        table.cell(1, 0).text = "Python"
        table.cell(1, 1).text = "Expert"

    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


# ----------------- PARSER TESTS -----------------

def test_parse_valid_pdf():
    pdf_bytes = create_sample_pdf()
    parsed = parse_resume_bytes(pdf_bytes, "resume.pdf")
    
    assert parsed.file_type == "pdf"
    assert parsed.page_count >= 1
    assert parsed.is_scanned_pdf is False
    assert parsed.ocr_recommended is False
    assert parsed.contact_info.email == "jane.doe@example.com"
    assert parsed.contact_info.phone == "(555) 123-4567"
    assert "linkedin.com/in/janedoe" in parsed.contact_info.linkedin
    assert "github.com/janedoe" in parsed.contact_info.github
    
    # Check sections extracted
    section_names = [s.name for s in parsed.sections]
    assert "summary" in section_names
    assert "experience" in section_names
    assert "skills" in section_names
    assert "education" in section_names


def test_detect_scanned_pdf():
    scanned_bytes = create_sample_pdf(is_scanned=True)
    parsed = parse_resume_bytes(scanned_bytes, "scanned_resume.pdf")
    
    assert parsed.is_scanned_pdf is True
    assert parsed.ocr_recommended is True
    assert any("OCR" in w or "Scanned" in w for w in parsed.warnings)


def test_parse_valid_docx():
    docx_bytes = create_sample_docx(include_table=True)
    parsed = parse_resume_bytes(docx_bytes, "resume.docx")
    
    assert parsed.file_type == "docx"
    assert parsed.contact_info.email == "jane.doe@example.com"
    assert parsed.is_scanned_pdf is False
    assert any("table" in w.lower() for w in parsed.warnings)


def test_empty_file_rejection():
    with pytest.raises(EmptyFileError):
        parse_resume_bytes(b"", "resume.pdf")


def test_invalid_extension():
    with pytest.raises(InvalidFileType):
        parse_resume_bytes(b"some content", "script.py")


def test_invalid_signature():
    with pytest.raises(InvalidFileType):
        parse_resume_bytes(b"THIS IS NOT A PDF", "fake.pdf")


def test_oversized_file():
    huge_bytes = b"%PDF-" + b"0" * (11 * 1024 * 1024)
    with pytest.raises(OversizedFileError):
        parse_resume_bytes(huge_bytes, "large.pdf")


# ----------------- API INTEGRATION TESTS -----------------

def test_api_health():
    response = client.get("/api/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "healthy"
    assert "app_name" in data


def test_api_parse_pdf():
    pdf_bytes = create_sample_pdf()
    files = {"file": ("resume.pdf", pdf_bytes, "application/pdf")}
    response = client.post("/api/resumes/parse", files=files)
    
    assert response.status_code == 200
    payload = response.json()
    assert payload["success"] is True
    assert payload["data"]["contact_info"]["email"] == "jane.doe@example.com"


def test_api_parse_docx():
    docx_bytes = create_sample_docx()
    files = {"file": ("resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = client.post("/api/resumes/parse", files=files)
    
    assert response.status_code == 200
    payload = response.json()
    assert payload["success"] is True
    assert payload["data"]["file_type"] == "docx"


def test_api_parse_empty_file():
    files = {"file": ("empty.pdf", b"", "application/pdf")}
    response = client.post("/api/resumes/parse", files=files)
    assert response.status_code == 400


def test_api_parse_unsupported_file():
    files = {"file": ("notes.txt", b"plain text", "text/plain")}
    response = client.post("/api/resumes/parse", files=files)
    assert response.status_code == 415
